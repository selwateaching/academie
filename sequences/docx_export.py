"""Export Word (.docx) d'une séquence : une fiche pédagogique modifiable par séance."""
import io
from urllib.parse import quote
import re

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

GREEN, ORANGE, RED, TEAL, GREY = "16A34A", "D97706", "DC2626", "0F766E", "6B7280"
LVL = {"easy": ("Facile", GREEN), "medium": ("Intermédiaire", ORANGE), "hard": ("Difficile", RED)}
PRODS = {"orale": "Production orale", "ecrite": "Production écrite", "application": "Exercices d'application",
         "debat": "Débat", "projet": "Projet", "lecture": "Lecture"}
HINT = "………………………………………………"
EARLY = re.compile(r"^(préscolaire|[123]ap|cp|ce1|ce2|maternelle)", re.I)
ARCH = re.compile(r"[؀-ۿ]")


def t(x, n=2000):
    return str(x or "").strip()[:n]


PROV_OWN = "Cours élaboré à partir du texte fourni par l'enseignant."
PROV = "Ce cours a été rédigé par une IA à partir de connaissances générales : aucune source n'a été consultée en direct. Les liens ci-dessous ouvrent des recherches pour vérifier chaque notion."
AR = {"À compléter par l'enseignant(e) directement dans Word : les cases « …… » sont modifiables (compétences, support, durée, enseignant). Les autres contenus sont générés et modifiables aussi.": 'يملأ الأستاذ(ة) الخانات « …… » مباشرة في Word (الكفاءات، السند، المدة، الأستاذ). باقي المحتويات مولَّدة وقابلة للتعديل أيضا.',
      'Présenter et faire expliquer le vocabulaire clé (tableau ci-dessous).': 'عرض المفردات الأساسية وشرحها (الجدول أدناه).',
      'Donner la consigne de production et les critères de réussite.': 'إعطاء تعليمة الإنتاج ومعايير النجاح.',
      'Évaluer les acquis (questions de vérification).': 'تقويم المكتسبات (أسئلة التحقق).',
      'Répètent et expliquent les mots.': 'يردّدون الكلمات ويشرحونها.',
      'Réalisent les exercices.': 'ينجزون التمارين.',
      'Réalisent la production.': 'ينجزون الإنتاج.',
      'Répondent aux questions.': 'يجيبون عن الأسئلة.',
      'FICHE PÉDAGOGIQUE': 'بطاقة بيداغوجية',
      'Niveau : ': 'المستوى : ',
      'Enseignant(e) : ': 'الأستاذ(ة) : ',
      'Compétence globale : ': 'الكفاءة الشاملة : ',
      'Compétence terminale : ': 'الكفاءة الختامية : ',
      'Séquence : ': 'المقطع : ',
      'Séance : ': 'الحصة : ',
      'Activité : ': 'النشاط : ',
      'Support : ': 'السند : ',
      'Titre : ': 'العنوان : ',
      'Durée : ': 'المدة : ',
      "🎯  Objectifs d'apprentissage": '🎯  أهداف التعلم',
      "L'apprenant sera capable de :": 'يكون المتعلم قادرا على :',
      'Déroulement de la séance': 'سير الحصة',
      'Phases': 'المراحل',
      "Tâches de l'enseignant": 'مهام الأستاذ',
      "Réponses (tâches) de l'élève": 'إجابات التلميذ',
      'Mise en situation': 'وضعية الانطلاق',
      'Construction du savoir': 'بناء التعلمات',
      'Vocabulaire': 'المفردات',
      'Application': 'التطبيق',
      'Production': 'الإنتاج',
      'Évaluation': 'التقويم',
      'Vocabulaire clé': 'المفردات الأساسية',
      'Exercices': 'تمارين',
      'Vérification des acquis': 'تقويم المكتسبات',
      'Ma production :': 'إنتاجي :',
      'Consigne': 'التعليمة',
      'Pour réussir la production': 'لأنجح في إنتاجي',
      "Corrigés (réservé à l'enseignant)": 'الإجابات النموذجية (خاصة بالأستاذ)',
      "Notes de l'enseignant": 'ملاحظات الأستاذ',
      'Critère': 'المعيار',
      'Points': 'النقاط',
      'Exemple attendu : ': 'مثال منتظر : ',
      'Note : ': 'ملاحظة : ',
      'Facile': 'سهل',
      'Intermédiaire': 'متوسط',
      'Difficile': 'صعب'}

AR.update({"Corrigés : voir la dernière partie du document (réservée à l'enseignant).": 'الإجابات النموذجية: انظر الجزء الأخير من الوثيقة (خاص بالأستاذ).', PROV_OWN: 'أُعدّ هذا الدرس انطلاقا من النص الذي قدّمه الأستاذ.', 'Provenance du contenu': 'مصدر المحتوى', 'Références à consulter': 'مراجع يُنصح بها', PROV: 'كُتب هذا الدرس بواسطة الذكاء الاصطناعي اعتمادا على معارف عامة، دون الاطلاع على أي مصدر مباشر. تؤدي الروابط أدناه إلى عمليات بحث للتحقق من كل مفهوم.', 'Texte support': 'النص السند', 'Découverte du support': 'اكتشاف السند', 'Faire lire le texte support (silencieusement puis à voix haute), vérifier la compréhension.': 'قراءة النص السند (صامتة ثم جهرية) والتحقق من الفهم.', 'Lisent le texte et répondent oralement.': 'يقرؤون النص ويجيبون شفهيا.', 'Sources': 'المصادر', 'consulté le': 'تاريخ الاطلاع'})


class Builder:
    def __init__(self, series):
        self.s = series
        lang = t(series.get("language")).lower()
        txt = t(series.get("title")) + t(series.get("topic"))
        self.rtl = lang.startswith(("arab", "hébr", "persan")) or len(ARCH.findall(txt)) >= 3
        self.big = bool(EARLY.match(t(series.get("level"))))
        self.size = 16 if self.big else 11
        self.doc = Document()
        sec = self.doc.sections[0]
        sec.left_margin = sec.right_margin = Cm(1.8)
        sec.top_margin = sec.bottom_margin = Cm(1.6)
        st = self.doc.styles["Normal"]
        st.font.name = "Arial"
        st.font.size = Pt(self.size)
        st.element.rPr.rFonts.set(qn("w:cs"), "Arial")
        st.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")

    def L(self, fr):
        return AR.get(fr, fr) if self.rtl else fr

    # ---- primitives
    def _rtl_par(self, p):
        if self.rtl:
            pPr = p._p.get_or_add_pPr()
            pPr.append(OxmlElement("w:bidi"))
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    def run(self, p, text, bold=False, size=None, color=None, italic=False):
        r = p.add_run(text)
        r.bold, r.italic = bold, italic
        r.font.size = Pt(size or self.size)
        if color:
            r.font.color.rgb = RGBColor.from_string(color)
        if self.rtl:
            rPr = r._r.get_or_add_rPr()
            rPr.append(OxmlElement("w:rtl"))
            for tag in ("w:bCs",) if bold else ():
                rPr.append(OxmlElement(tag))
            szcs = OxmlElement("w:szCs")
            szcs.set(qn("w:val"), str(int((size or self.size) * 2)))
            rPr.append(szcs)
        return r

    def par(self, host, text="", bold=False, size=None, color=None, italic=False, align=None, after=4, style=None):
        p = host.add_paragraph(style=style) if style else host.add_paragraph()
        if text:
            self.run(p, text, bold, size, color, italic)
        pf = p.paragraph_format
        pf.space_after, pf.space_before = Pt(after), Pt(0)
        pf.line_spacing = 1.8 if self.big else 1.25
        self._rtl_par(p)
        if align == "c":
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        return p

    def cell_text(self, cell, text, bold=False, size=None, color=None, italic=False, first=True):
        p = cell.paragraphs[0] if first and not cell.paragraphs[0].text else cell.add_paragraph()
        if text:
            self.run(p, text, bold, size, color, italic)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.7 if self.big else 1.2
        self._rtl_par(p)
        return p

    @staticmethod
    def shade(cell, color):
        tcPr = cell._tc.get_or_add_tcPr()
        sh = OxmlElement("w:shd")
        sh.set(qn("w:val"), "clear")
        sh.set(qn("w:color"), "auto")
        sh.set(qn("w:fill"), color)
        tcPr.append(sh)

    def table(self, rows, cols, widths=None):
        tb = self.doc.add_table(rows=rows, cols=cols)
        tb.style = "Table Grid"
        tb.alignment = WD_TABLE_ALIGNMENT.CENTER
        if self.rtl:
            tb._tbl.tblPr.append(OxmlElement("w:bidiVisual"))
        if widths:
            for row in tb.rows:
                for c, w in zip(row.cells, widths):
                    c.width = Cm(w)
        return tb

    def heading(self, text, color=TEAL):
        tb = self.table(1, 1)
        c = tb.rows[0].cells[0]
        self.shade(c, "E6F4F1")
        self.cell_text(c, text, bold=True, size=self.size + 2, color=color)
        self.par(self.doc, "", after=2)

    def lines(self, n, host=None):
        """Lignes de cahier : un tableau sans cadre dont chaque rangée a un trait en bas (Word fusionnerait des paragraphes bordés)."""
        tb = self.doc.add_table(rows=n, cols=1)
        tbl = tb._tbl
        pr = tbl.tblPr
        borders = OxmlElement("w:tblBorders")
        for side in ("top", "left", "bottom", "right", "insideV"):
            e = OxmlElement("w:" + side)
            e.set(qn("w:val"), "nil")
            borders.append(e)
        for side in ("insideH", "bottom"):
            e = OxmlElement("w:" + side)
            for k, v in (("val", "single"), ("sz", "6"), ("space", "0"), ("color", "7DA7D9")):
                e.set(qn("w:" + k), v)
            borders.append(e)
        pr.append(borders)
        if self.rtl:
            pr.append(OxmlElement("w:bidiVisual"))
        h = 1.25 if self.big else 0.95
        for row in tb.rows:
            trPr = row._tr.get_or_add_trPr()
            hh = OxmlElement("w:trHeight")
            hh.set(qn("w:val"), str(int(h * 567)))
            hh.set(qn("w:hRule"), "exact")
            trPr.append(hh)
            cant = OxmlElement("w:cantSplit")
            trPr.append(cant)
            c = row.cells[0]
            c.width = Cm(18)
            c.paragraphs[0].paragraph_format.space_after = Pt(0)
        self.par(self.doc, "", after=2)

    def link(self, p, text, url):
        rid = p.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
        h = OxmlElement("w:hyperlink")
        h.set(qn("r:id"), rid)
        r = OxmlElement("w:r")
        rPr = OxmlElement("w:rPr")
        c = OxmlElement("w:color")
        c.set(qn("w:val"), "0563C1")
        u = OxmlElement("w:u")
        u.set(qn("w:val"), "single")
        rPr.append(c)
        rPr.append(u)
        r.append(rPr)
        tx = OxmlElement("w:t")
        tx.text = text
        r.append(tx)
        h.append(r)
        p._p.append(h)

    def lvl_tag(self, p, level):
        if level in LVL:
            self.run(p, "● " + self.L(LVL[level][0]) + "  ", True, self.size - 1, LVL[level][1])

    # ---- sections
    def fiche(self, i, lesson):
        n = len(self.s["lessons"])
        sl = [x for x in lesson.get("slides") or [] if isinstance(x, dict)]
        title_sl = next((x for x in sl if x.get("layout") == "title"), {})
        content = [x for x in sl if x.get("layout") in ("text_image", "list", "steps", "timeline", "quote", "video")]
        vocab = [v for x in sl if x.get("layout") == "vocab" for v in x.get("vocab") or []]
        acts = [x for x in sl if x.get("layout") == "activity" and isinstance(x.get("q"), dict)]
        exo = [x for x in acts if not t(x.get("title")).lower().startswith("vérification")]
        quiz = [x for x in acts if x not in exo]
        sups = [x for x in sl if x.get("layout") == "support" and t(x.get("text"))]
        sup = ({"title": re.sub(r"\s*\(\d+/\d+\)$", "", t(sups[0].get("title"))), "text": "\n".join(t(x.get("text"), 3000) for x in sups)} if sups else None)
        prods = [x for x in sl if x.get("layout") == "production" and isinstance(x.get("prod"), dict)]

        tb = self.table(1, 1)
        c = tb.rows[0].cells[0]
        self.shade(c, "DFF1EE")
        p = self.cell_text(c, self.L("FICHE PÉDAGOGIQUE"), bold=True, size=self.size + 8, color="1D4ED8")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        self.par(self.doc, self.L("À compléter par l'enseignant(e) directement dans Word : les cases « …… » sont modifiables (compétences, support, durée, enseignant). Les autres contenus sont générés et modifiables aussi."),
                 italic=True, size=max(self.size - 3, 8), color=GREY, after=6)

        h = self.table(6, 2, [9, 9])
        cells = [
            ((self.L("Niveau : "), t(self.s.get("level"))), (self.L("Enseignant(e) : "), HINT)),
            ((self.L("Compétence globale : "), HINT + HINT + HINT), None),
            ((self.L("Compétence terminale : "), HINT + HINT + HINT), None),
            ((self.L("Séquence : "), t(self.s.get("title"))), (self.L("Séance : "), f"{i + 1} / {n}")),
            ((self.L("Activité : "), HINT), (self.L("Support : "), t(sup.get("title"), 120) if sup else HINT)),
            ((self.L("Titre : "), t(lesson.get("title"))), (self.L("Durée : "), "1 h")),
        ]
        for r, (a, b) in enumerate(cells):
            row = h.rows[r]
            if b is None:
                m = row.cells[0].merge(row.cells[1])
                self.cell_text(m, a[0], True, color=TEAL)
                self.run(m.paragraphs[0], a[1], italic=True, color=GREY)
                continue
            for cell, (lab, val) in zip(row.cells, (a, b)):
                self.cell_text(cell, lab, True, color=TEAL)
                self.run(cell.paragraphs[0], val, italic=(val == HINT), color=GREY if val == HINT else None)
        self.par(self.doc, "", after=4)

        self.heading(self.L("🎯  Objectifs d'apprentissage"))
        obj = [t(b, 400) for b in title_sl.get("bullets") or [] if b]
        self.par(self.doc, self.L("L'apprenant sera capable de :"), bold=True)
        for b in obj or ["……………………………………………………………"]:
            self.par(self.doc, "•  " + b, after=2)
        self.par(self.doc, "", after=4)

        if sup:
            self.heading(self.L("Texte support") + (" : " + t(sup.get("title"), 160) if sup.get("title") else ""))
            tbs = self.table(1, 1)
            self.shade(tbs.rows[0].cells[0], "FFFBEA")
            for k, para in enumerate(t(sup.get("text"), 20000).split("\n")):
                if para.strip():
                    self.cell_text(tbs.rows[0].cells[0], para.strip(), size=self.size + (2 if self.big else 1), first=(k == 0))
            self.par(self.doc, "", after=4)

        self.heading(self.L("Déroulement de la séance"))
        phases = []
        if sup:
            phases.append((self.L("Découverte du support"), "support"))
        if content:
            phases.append((self.L("Mise en situation"), content[:1]))
            if content[1:]:
                phases.append((self.L("Construction du savoir"), content[1:]))
        if vocab:
            phases.append((self.L("Vocabulaire"), "vocab"))
        if exo:
            phases.append((self.L("Application"), "exo"))
        if prods:
            phases.append((self.L("Production"), "prod"))
        if quiz:
            phases.append((self.L("Évaluation"), "quiz"))
        tb = self.table(1 + len(phases), 3, [3.6, 7.6, 6.8])
        for cell, txt in zip(tb.rows[0].cells, (self.L("Phases"), self.L("Tâches de l'enseignant"), self.L("Réponses (tâches) de l'élève"))):
            self.shade(cell, "FFF4C2")
            self.cell_text(cell, txt, True, color=TEAL)
        for r, (name, what) in enumerate(phases, start=1):
            a, b, c = tb.rows[r].cells
            self.cell_text(a, name, True, color=TEAL)
            if isinstance(what, list):
                for sx in what:
                    self.slide_text(b, sx)
                self.cell_text(c, "……………………………", italic=True, color=GREY)
            elif what == "support":
                self.cell_text(b, self.L("Faire lire le texte support (silencieusement puis à voix haute), vérifier la compréhension."))
                self.cell_text(c, self.L("Lisent le texte et répondent oralement."), italic=True, color=GREY)
            elif what == "vocab":
                self.cell_text(b, self.L("Présenter et faire expliquer le vocabulaire clé (tableau ci-dessous)."))
                self.cell_text(c, self.L("Répètent et expliquent les mots."), italic=True, color=GREY)
            elif what == "exo":
                self.cell_text(b, (f"Faire réaliser les {len(exo)} exercices." if not self.rtl else f"إنجاز التمارين ({len(exo)})."))
                self.cell_text(c, self.L("Réalisent les exercices."), italic=True, color=GREY)
            elif what == "prod":
                self.cell_text(b, self.L("Donner la consigne de production et les critères de réussite."))
                self.cell_text(c, self.L("Réalisent la production."), italic=True, color=GREY)
            else:
                self.cell_text(b, self.L("Évaluer les acquis (questions de vérification)."))
                self.cell_text(c, self.L("Répondent aux questions."), italic=True, color=GREY)
        self.par(self.doc, "", after=4)

        if vocab:
            self.heading(self.L("Vocabulaire clé"))
            tb = self.table(len(vocab), 2, [5, 13])
            for r, v in enumerate(vocab):
                self.cell_text(tb.rows[r].cells[0], t(v.get("term"), 100), True, color=TEAL)
                self.cell_text(tb.rows[r].cells[1], t(v.get("definition"), 400))
            self.par(self.doc, "", after=4)

        if exo:
            self.heading(self.L("Exercices"))
            self.par(self.doc, self.L("Corrigés : voir la dernière partie du document (réservée à l'enseignant)."), italic=True,
                     size=max(self.size - 2, 8), color=GREY, after=4)
            for k, x in enumerate(exo, 1):
                self.exercise(k, x["q"])

        for x in prods:
            self.production(x["prod"])

        if quiz:
            self.heading(self.L("Vérification des acquis"))
            for k, x in enumerate(quiz, 1):
                self.exercise(k, x["q"], lvl=False)

        srcs = [v for x in sl if x.get("layout") == "sources" for v in x.get("sources") or [] if isinstance(v, dict) and (v.get("url") or v.get("title"))]
        prov = next((x.get("prov") for x in sl if x.get("layout") == "sources" and isinstance(x.get("prov"), dict)), None)
        if srcs or prov:
            self.heading(self.L("Sources"))
            for v in srcs:
                p = self.par(self.doc, "", after=3)
                if v.get("url"):
                    self.link(p, t(v.get("title") or v.get("site") or v["url"], 200), t(v["url"], 500))
                    extra = ", ".join(x for x in (t(v.get("site"), 80), (self.L("consulté le") + " " + t(v.get("accessed"), 20)) if v.get("accessed") else "") if x)
                    if extra:
                        self.run(p, f"  ({extra})", size=max(self.size - 2, 8), color=GREY)
                else:
                    self.run(p, t(v.get("title"), 300))
            if prov:
                self.par(self.doc, self.L("Provenance du contenu"), bold=True, color=TEAL)
                self.par(self.doc, self.L(PROV_OWN if prov.get("src") else PROV), italic=True, size=max(self.size - 1, 9), color=GREY, after=3)
                if prov.get("works"):
                    self.par(self.doc, self.L("Références à consulter") + " : " + " · ".join(t(w, 120) for w in prov["works"]), after=3)
                wl = {"arab": "ar", "angl": "en", "espag": "es", "allem": "de"}.get(t(self.s.get("language")).lower()[:5].rstrip("e"), "fr")
                for k in prov.get("keywords") or []:
                    p = self.par(self.doc, "", after=2)
                    self.run(p, "•  " + t(k, 80) + "  ")
                    self.link(p, "Wikipédia", f"https://{wl}.wikipedia.org/w/index.php?search={quote(t(k, 80))}")
                    self.run(p, "  ")
                    self.link(p, "Google Scholar", f"https://scholar.google.com/scholar?q={quote(t(k, 80))}")

        self.corriges(exo, quiz, prods, sl)

    def slide_text(self, cell, sx):
        if sx.get("title"):
            self.cell_text(cell, t(sx["title"], 200), True)
        if sx.get("layout") == "quote":
            self.cell_text(cell, "« " + t(sx.get("quote"), 400) + " » — " + t(sx.get("author"), 100), italic=True)
        for it in sx.get("items") or []:
            if isinstance(it, dict):
                self.cell_text(cell, "•  " + t(it.get("t"), 200) + (" : " + t(it.get("x"), 400) if it.get("x") else ""))
        for b in sx.get("bullets") or []:
            self.cell_text(cell, "•  " + t(b, 500))
        if sx.get("notes"):
            self.cell_text(cell, self.L("Note : ") + t(sx["notes"], 700), italic=True, size=max(self.size - 2, 8), color=GREY)

    def exercise(self, k, q, lvl=True):
        p = self.par(self.doc, "", after=2)
        self.run(p, f"{k}.  ", True, color=TEAL)
        if lvl:
            self.lvl_tag(p, t(q.get("level")))
        self.run(p, t(q.get("q"), 700), True)
        for j, c in enumerate(q.get("choices") or []):
            self.par(self.doc, f"     {'abcdef'[j % 6]})  {t(c, 300)}", after=1)
        if not q.get("choices"):
            self.lines(2)
        self.par(self.doc, "", after=4)

    def production(self, pr):
        kind = PRODS.get(pr.get("kind"), self.L("Production"))
        lv = LVL.get(pr.get("level"))
        self.heading(("✍️  " if pr.get("kind") != "orale" else "🎤  ") + (t(pr.get("title")) or kind) + (f" — {self.L(lv[0])}" if lv else ""),
                     color=lv[1] if lv else TEAL)
        tb = self.table(1, 2, [11, 7])
        a, b = tb.rows[0].cells
        self.shade(a, "FFF1E6")
        self.shade(b, "E7F7EC")
        self.cell_text(a, self.L("Consigne") + (f" · {t(pr.get('duration'), 60)}" if pr.get("duration") else ""), True, color="C2410C")
        self.cell_text(a, t(pr.get("task"), 1500), first=False)
        self.cell_text(b, self.L("Pour réussir la production"), True, color=GREEN)
        for c in pr.get("criteria") or []:
            self.cell_text(b, "•  " + t(c, 240), first=False)
        self.par(self.doc, "", after=4)
        self.par(self.doc, self.L("Ma production :"), bold=True, color=TEAL)
        self.lines(12 if not self.big else 9)
        self.par(self.doc, "", after=4)

    def corriges(self, exo, quiz, prods, sl):
        self.doc.add_page_break()
        self.heading(self.L("Corrigés (réservé à l'enseignant)"), color=RED)
        for name, xs in ((self.L("Exercices"), exo), (self.L("Vérification des acquis"), quiz)):
            if not xs:
                continue
            self.par(self.doc, name, bold=True, color=TEAL)
            for k, x in enumerate(xs, 1):
                q = x["q"]
                ch = q.get("choices") or []
                a = (f"{'abcdef'[int(q.get('answer') or 0) % 6]}) {t(ch[int(q.get('answer') or 0)] if ch else '', 300)}"
                     if ch else t(q.get("open"), 900))
                p = self.par(self.doc, "", after=3)
                self.run(p, f"{k}.  ", True)
                if name == self.L("Exercices"):
                    self.lvl_tag(p, t(q.get("level")))
                self.run(p, a)
                if q.get("why"):
                    self.run(p, "  — " + t(q["why"], 700), italic=True, color=GREY)
        for x in prods:
            pr = x["prod"]
            self.par(self.doc, "Production : " + (t(pr.get("title")) or PRODS.get(pr.get("kind"), "")), bold=True, color=TEAL)
            if pr.get("example"):
                self.par(self.doc, self.L("Exemple attendu : ") + t(pr["example"], 1500), after=3)
            if pr.get("rubric"):
                tb = self.table(1 + len(pr["rubric"]), 2, [14, 4])
                tb.rows[0].cells[0].text, tb.rows[0].cells[1].text = self.L("Critère"), self.L("Points")
                for r, z in enumerate(pr["rubric"], 1):
                    tb.rows[r].cells[0].text, tb.rows[r].cells[1].text = t(z.get("criterion"), 200), str(z.get("points", ""))
        notes = [(t(x.get("title")), t(x.get("notes"), 900)) for x in sl if x.get("notes") and x.get("layout") not in ("production",)]
        if notes:
            self.par(self.doc, self.L("Notes de l'enseignant"), bold=True, color=TEAL)
            for ti, nt in notes:
                self.par(self.doc, f"{ti} : {nt}", size=max(self.size - 1, 9), after=2)

    def build(self):
        lessons = [x for x in self.s.get("lessons") or [] if isinstance(x, dict) and x.get("slides")]
        self.s["lessons"] = lessons
        for i, l in enumerate(lessons):
            if i:
                self.doc.add_page_break()
            self.fiche(i, l)
        bio = io.BytesIO()
        self.doc.save(bio)
        return bio.getvalue()


def build_docx(series):
    return Builder(series).build()
