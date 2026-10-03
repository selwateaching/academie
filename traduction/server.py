"""Traducteur Association — serveur (Render ou autre).

  /               la page (static/index.html)
  /api/translate  traduit un texte avec Claude (clé gardée côté serveur)
  /api/room/...   salles : relie deux téléphones (bénévole / personne accueillie)
  /api/document   traduit la photo d'un document (courrier, formulaire…)
  /api/transcribe transcrit un fichier audio reçu (WhatsApp, etc.) avec Whisper

Variables d'environnement :
  ANTHROPIC_API_KEY   obligatoire pour traduire
  OPENAI_API_KEY      facultatif : seulement pour importer des fichiers audio
  TRAD_ACCESS_CODE    facultatif : code d'accès demandé dans la page
  TRAD_MODEL          facultatif : modèle Claude (défaut claude-sonnet-5-5)
  TRAD_PER_HOUR       facultatif : requêtes max par heure et par IP (défaut 600)
"""
import base64
import io
import json
import hashlib
import hmac
import os
import sqlite3
import uuid
from contextlib import contextmanager
import secrets
import threading
import time
from collections import defaultdict, deque

import anthropic
import requests
import segno
from flask import Flask, Response, jsonify, request, send_from_directory

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
STATIC_DIR = os.path.join(BASE_DIR, "static")

MODEL = os.environ.get("TRAD_MODEL", "claude-sonnet-5-5")
ACCESS_CODE = os.environ.get("TRAD_ACCESS_CODE", "").strip() or ""
PER_HOUR = int(os.environ.get("TRAD_PER_HOUR", "600"))
MAX_TEXT = 4000

app = Flask(__name__, static_folder=None)
app.config["MAX_CONTENT_LENGTH"] = 25 * 1024 * 1024  # limite Whisper

_log = defaultdict(deque)
_lock = threading.Lock()
_client = None


def get_client():
    global _client
    if _client is None:
        _client = anthropic.Anthropic(timeout=60.0)
    return _client


def error(message, status):
    return jsonify({"error": message}), status


def check_access():
    """Renvoie une réponse d'erreur, ou None si la requête est autorisée."""
    if ACCESS_CODE:
        given = request.headers.get("X-Access-Code", "")
        ip0 = (request.headers.get("X-Forwarded-For", "").split(",")[0].strip()) or (request.remote_addr or "?")
        now0 = time.time()
        with _lock:
            fails = _fail_log[ip0]
            while fails and now0 - fails[0] > 900:
                fails.popleft()
            if len(fails) >= 10:
                return error("Trop d'essais de code. Réessayez dans 15 minutes.", 429)
        if not hmac.compare_digest(given.encode(), ACCESS_CODE.encode()):
            with _lock:
                _fail_log[ip0].append(now0)
            return error("Code d'accès incorrect.", 401)
    fwd = request.headers.get("X-Forwarded-For", "")
    ip = fwd.split(",")[0].strip() if fwd else (request.remote_addr or "?")
    now = time.time()
    with _lock:
        q = _log[ip]
        while q and now - q[0] > 3600:
            q.popleft()
        if len(q) >= PER_HOUR:
            return error("Trop de requêtes, réessayez plus tard.", 429)
        q.append(now)
    return None


@app.get("/")
def index():
    with open(os.path.join(STATIC_DIR, "index.html"), encoding="utf-8") as f:
        html = f.read()
    html = html.replace("__NEEDS_CODE__", "true" if ACCESS_CODE else "false")
    return Response(html, mimetype="text/html", headers={"Cache-Control": "no-cache"})


@app.get("/sw.js")
def service_worker():
    return send_from_directory(STATIC_DIR, "sw.js", mimetype="text/javascript", max_age=0)


@app.get("/manifest.json")
def manifest():
    return send_from_directory(STATIC_DIR, "manifest.json", mimetype="application/manifest+json")


ICONS = {"icon-512.png", "icon-192.png", "apple-touch-icon.png", "favicon-48.png", "aide-grand.png"}


@app.get("/<name>.png")
def icon_png(name):
    if name + ".png" not in ICONS:
        return error("Introuvable.", 404)
    return send_from_directory(STATIC_DIR, name + ".png", mimetype="image/png", max_age=86400)


@app.get("/icon.svg")
def icon():
    return send_from_directory(STATIC_DIR, "icon.svg", mimetype="image/svg+xml")


@app.get("/healthz")
def healthz():
    return {
        "ok": True,
        "translate": bool(os.environ.get("ANTHROPIC_API_KEY")),
        "transcribe": bool(os.environ.get("OPENAI_API_KEY")),
        "dossiers": bool(ACCESS_CODE and DOSSIER_KEY),
        "TRAD_ACCESS_CODE": bool(ACCESS_CODE),
        "DOSSIER_KEY": bool(DOSSIER_KEY),
    }


def translate_text(text, src, dst, context=""):
    """Traduit avec Claude. Lève anthropic.APIError en cas de problème."""
    system = (
        "Tu es un interprète professionnel dans une association qui aide des personnes étrangères "
        "(démarches administratives, santé, logement, école, vie quotidienne). "
        f"Traduis fidèlement le message de l'utilisateur {'depuis : ' + src + ' ' if src else ''}"
        f"vers : {dst}. Le texte vient d'une reconnaissance vocale et peut contenir des erreurs "
        "ou des hésitations : corrige-les avec bon sens sans rien inventer. "
        "Garde le ton et le niveau de langue, avec des mots simples. "
        "Pour l'arabe algérien, tunisien ou marocain (darija) : le texte reçu mélange souvent darija "
        "et mots français, écrits en lettres arabes ou latines (arabizi) ; comprends-les, et traduis vers "
        "le français naturel. Quand la cible est une darija, écris en lettres arabes, dans le dialecte "
        "demandé (pas en arabe littéraire), avec des mots simples et courants, en gardant tels quels "
        "les termes administratifs français usuels (CAF, préfecture, rendez-vous…). "
        "Réponds UNIQUEMENT par la traduction, sans commentaire, sans guillemets. "
        "Si le message est incompréhensible, réponds exactement : [incompréhensible]"
    )
    if context:
        system += f"\n\nDébut de la conversation, pour le contexte :\n{context}"
    msg = get_client().messages.create(
        model=MODEL,
        max_tokens=1500,
        system=system,
        messages=[{"role": "user", "content": text}],
    )
    return "".join(b.text for b in msg.content if b.type == "text").strip()


@app.post("/api/translate")
def translate():
    denied = check_access()
    if denied:
        return denied
    if not os.environ.get("ANTHROPIC_API_KEY"):
        return error("Le serveur n'a pas encore de clé API Anthropic.", 503)
    data = request.get_json(silent=True) or {}
    text = str(data.get("text", "")).strip()[:MAX_TEXT]
    src = str(data.get("from", "")).strip()[:40]
    dst = str(data.get("to", "")).strip()[:40]
    if not text or not dst:
        return error("Texte ou langue manquants.", 400)
    context = str(data.get("context", "")).strip()[:1500]
    try:
        return jsonify({"translation": translate_text(text, src, dst, context)})
    except anthropic.APIError as e:
        return error(f"Erreur de traduction : {getattr(e, 'message', e)}", 502)


# ------------------------------------------------------------ salles (2 téléphones)
# Une salle relie le téléphone du bénévole (français) et celui de la personne
# (sa langue). Gardées en mémoire (1 seul worker gunicorn) et effacées après 3 h.
ROOM_TTL = 3 * 3600
ROOM_MAX_MSGS = 400
ROOM_MAX_ROOMS = 200
ALPHABET = "ABCDEFGHJKLMNPQRSTUVWXYZ23456789"
rooms = {}
rooms_lock = threading.Lock()


def get_room(code):
    now = time.time()
    with rooms_lock:
        for c in [c for c, r in rooms.items() if now - r["created"] > ROOM_TTL]:
            del rooms[c]
        return rooms.get(str(code).upper())


@app.post("/api/room")
def room_create():
    denied = check_access()
    if denied:
        return denied
    if not os.environ.get("ANTHROPIC_API_KEY"):
        return error("Le serveur n'a pas encore de clé API Anthropic.", 503)
    get_room("")  # nettoie les salles expirées
    with rooms_lock:
        if len(rooms) >= ROOM_MAX_ROOMS:
            return error("Trop de sessions ouvertes, réessayez plus tard.", 429)
        code = "".join(secrets.choice(ALPHABET) for _ in range(6))
        rooms[code] = {"created": time.time(), "guest_code": "", "guest_name": "", "msgs": []}
    return jsonify({"code": code})


@app.post("/api/room/<code>/join")
def room_join(code):
    r = get_room(code)
    if not r:
        return error("Session introuvable ou expirée.", 404)
    d = request.get_json(silent=True) or {}
    r["guest_code"] = str(d.get("lang_code", ""))[:12]
    r["guest_name"] = str(d.get("lang_name", ""))[:60]
    if not r["guest_name"]:
        return error("Langue manquante.", 400)
    return jsonify({"ok": True})


@app.post("/api/room/<code>/say")
def room_say(code):
    r = get_room(code)
    if not r:
        return error("Session introuvable ou expirée.", 404)
    d = request.get_json(silent=True) or {}
    side = d.get("side")
    text = str(d.get("text", "")).strip()[:MAX_TEXT]
    if side not in ("host", "guest") or not text:
        return error("Message invalide.", 400)
    if not r["guest_name"]:
        return error("La personne n'a pas encore rejoint la session.", 409)
    if len(r["msgs"]) >= ROOM_MAX_MSGS:
        return error("Session pleine : créez-en une nouvelle.", 429)
    fr, other = "français", r["guest_name"]
    src, dst = (fr, other) if side == "host" else (other, fr)
    context = "\n".join(
        f"{'Association' if m['side'] == 'host' else 'Personne'}: {m['text']}" for m in r["msgs"][-6:]
    )
    try:
        tr = translate_text(text, src, dst, context)
    except anthropic.APIError as e:
        return error(f"Erreur de traduction : {getattr(e, 'message', e)}", 502)
    with rooms_lock:
        mid = (r["msgs"][-1]["id"] + 1) if r["msgs"] else 1
        r["msgs"].append({"id": mid, "side": side, "text": text, "tr": tr})
    return jsonify({"id": mid})


@app.delete("/api/room/<code>")
def room_delete(code):
    with rooms_lock:
        rooms.pop(str(code).upper(), None)
    return jsonify({"ok": True})


@app.get("/api/room/<code>")
def room_poll(code):
    r = get_room(code)
    if not r:
        return error("Session introuvable ou expirée.", 404)
    since = request.args.get("since", 0, type=int)
    return jsonify({
        "guest_code": r["guest_code"],
        "guest_name": r["guest_name"],
        "msgs": [m for m in r["msgs"] if m["id"] > since],
    })


PHRASES = [
    "Bonjour, bienvenue.",
    "Comment vous appelez-vous ?",
    "Avez-vous un rendez-vous ?",
    "Asseyez-vous, s'il vous plaît.",
    "Attendez ici, s'il vous plaît.",
    "Avez-vous une pièce d'identité ?",
    "Pouvez-vous remplir ce formulaire ?",
    "Signez ici, s'il vous plaît.",
    "Je ne comprends pas. Pouvez-vous répéter ?",
    "Parlez lentement, s'il vous plaît.",
    "Avez-vous compris ?",
    "Avez-vous des enfants ? Combien ?",
    "Avez-vous besoin d'un médecin ?",
    "Apportez vos papiers : passeport, justificatif de domicile, avis d'imposition.",
    "Revenez lundi à 10 heures.",
    "C'est gratuit.",
    "Je vais utiliser l'application pour vous traduire.",
    "Merci, au revoir.",
]
_phrase_cache = {}


@app.post("/api/phrases")
def phrases():
    """Phrases d'accueil traduites ; le navigateur les garde pour le hors connexion."""
    denied = check_access()
    if denied:
        return denied
    if not os.environ.get("ANTHROPIC_API_KEY"):
        return error("Le serveur n'a pas encore de clé API Anthropic.", 503)
    dst = str((request.get_json(silent=True) or {}).get("to", "")).strip()[:60]
    if not dst:
        return error("Langue manquante.", 400)
    if dst not in _phrase_cache:
        system = (
            "Tu es un interprète dans une association d'accueil de personnes étrangères. "
            f"Traduis chaque phrase française de la liste en {dst}, en mots simples et polis. "
            "Si la langue est une darija, écris en lettres arabes, dans le dialecte demandé. "
            "Réponds UNIQUEMENT par un tableau JSON de chaînes, dans le même ordre, de même longueur."
        )
        try:
            msg = get_client().messages.create(
                model=MODEL, max_tokens=4000, system=system,
                messages=[{"role": "user", "content": json.dumps(PHRASES, ensure_ascii=False)}],
            )
            text = "".join(b.text for b in msg.content if b.type == "text")
            out = json.loads(text[text.index("["): text.rindex("]") + 1])
        except anthropic.APIError as e:
            return error(f"Erreur de traduction : {getattr(e, 'message', e)}", 502)
        except ValueError:
            return error("Réponse de traduction illisible, réessayez.", 502)
        if not isinstance(out, list) or len(out) != len(PHRASES):
            return error("Réponse de traduction incomplète, réessayez.", 502)
        _phrase_cache[dst] = [str(x) for x in out]
    return jsonify({"phrases": [{"fr": f, "tr": t} for f, t in zip(PHRASES, _phrase_cache[dst])]})



# ------------------------------------------------------------ dossiers (suivi) et courriers
# Les dossiers contiennent des données personnelles : ils sont chiffrés (Fernet) avant d'être
# écrits dans SQLite, et ne sont accessibles qu'avec le code d'accès. Désactivés tant que
# TRAD_ACCESS_CODE et DOSSIER_KEY ne sont pas définis.
from cryptography.fernet import Fernet, InvalidToken

def read_setting(name):
    """Valeur d'un réglage : variable d'environnement, ou « Secret File » Render du même nom."""
    v = os.environ.get(name, "").strip()
    if v:
        return v
    for folder in ("/etc/secrets", BASE_DIR, os.getcwd()):
        try:
            with open(os.path.join(folder, name), encoding="utf-8") as f:
                v = f.read().strip()
            if v:
                return v
        except OSError:
            pass
    return ""


DOSSIER_KEY = read_setting("DOSSIER_KEY")
RETENTION_MONTHS = int(os.environ.get("DOSSIER_RETENTION_MONTHS", "0"))  # 0 = jamais de suppression automatique
DATA_DIR = os.environ.get("DATA_DIR") or ("/var/data" if os.path.isdir("/var/data") else os.path.join(BASE_DIR, "data"))
DB_PATH = os.path.join(DATA_DIR, "dossiers.db")
_db_lock = threading.Lock()
_fernet = None
_fail_log = defaultdict(deque)  # tentatives de code d'accès ratées, par IP


def fernet():
    global _fernet
    if _fernet is None:
        raw = hashlib.pbkdf2_hmac("sha256", DOSSIER_KEY.encode(), b"tarjama-dossiers-v1", 200_000)
        _fernet = Fernet(base64.urlsafe_b64encode(raw))
    return _fernet


@contextmanager
def db():
    os.makedirs(DATA_DIR, exist_ok=True)
    con = sqlite3.connect(DB_PATH, timeout=10)
    try:
        con.execute("CREATE TABLE IF NOT EXISTS dossiers (id TEXT PRIMARY KEY, updated REAL, data BLOB)")
        yield con
    finally:
        con.close()


def dossier_guard():
    missing = [n for n, v in (("TRAD_ACCESS_CODE", ACCESS_CODE), ("DOSSIER_KEY", DOSSIER_KEY)) if not v]
    if missing:
        return error("Dossiers non activés : il manque sur Render la variable " + " et ".join(missing) + ".", 503)
    return check_access()


def load_dossier(con, did):
    row = con.execute("SELECT data FROM dossiers WHERE id=?", (did,)).fetchone()
    if not row:
        return None
    try:
        d = json.loads(fernet().decrypt(row[0]))
    except InvalidToken:
        return None
    d["id"] = did
    return d


def save_dossier(con, d):
    did = d.pop("id")
    d["updated"] = time.time()
    con.execute("INSERT OR REPLACE INTO dossiers VALUES (?,?,?)",
                (did, d["updated"], fernet().encrypt(json.dumps(d, ensure_ascii=False).encode())))
    con.commit()
    d["id"] = did


def clip(v, n):
    return str(v or "").strip()[:n]


FIELDS = {"prenom": 60, "nom": 60, "langue_code": 12, "langue_nom": 60, "tel": 30,
          "statut": 20, "echeance": 10, "notes": 3000,
          "metier": 80, "niveau_fr": 20, "disponibilite": 80, "permis": 10, "entreprise": 100,
          "etape": 60, "titre_type": 60, "titre_expiration": 10, "consentement": 10}
DATE_FIELDS = ("echeance", "titre_expiration", "consentement")
with open(os.path.join(BASE_DIR, "parcours.json"), encoding="utf-8") as _f:
    PARCOURS = json.load(_f)
PIECE_KEYS = {p["key"] for p in PARCOURS["pieces"]}


def clean_dossier(d):
    """Normalise les champs d'un dossier (dates, statut, pièces)."""
    for k in DATE_FIELDS:
        d[k] = to_date(d.get(k))
    if d.get("statut") not in ("En cours", "En attente", "Clos"):
        d["statut"] = "En cours"
    d["pieces"] = [k for k in dict.fromkeys(d.get("pieces") or []) if k in PIECE_KEYS]
    return d


def summary(d):
    out = {k: d.get(k, "") for k in ("id", "prenom", "nom", "langue_nom", "statut", "echeance", "updated",
                                      "metier", "entreprise", "etape", "titre_expiration", "niveau_fr")}
    out["pieces_ok"] = len(d.get("pieces") or [])
    return out


@app.get("/api/parcours")
def parcours():
    denied = check_access()
    return denied or jsonify(PARCOURS)


@app.get("/api/dossiers")
def dossiers_list():
    denied = dossier_guard()
    if denied:
        return denied
    with _db_lock, db() as con:
        if RETENTION_MONTHS > 0:  # suppression automatique des dossiers inactifs
            con.execute("DELETE FROM dossiers WHERE updated < ?", (time.time() - RETENTION_MONTHS * 30 * 86400,))
            con.commit()
        ids = [r[0] for r in con.execute("SELECT id FROM dossiers")]
        out = [summary(d) for d in (load_dossier(con, i) for i in ids) if d]
    out.sort(key=lambda d: -d["updated"])
    return jsonify({"dossiers": out})


@app.post("/api/dossiers")
def dossier_create():
    denied = dossier_guard()
    if denied:
        return denied
    data = request.get_json(silent=True) or {}
    d = {k: clip(data.get(k), n) for k, n in FIELDS.items()}
    if not d["prenom"] and not d["nom"]:
        return error("Indiquez au moins un prénom ou un nom.", 400)
    d["pieces"] = [str(x)[:40] for x in (data.get("pieces") or [])][:40]
    clean_dossier(d)
    d["created"] = time.time()
    d["journal"] = []
    d["id"] = uuid.uuid4().hex[:12]
    with _db_lock, db() as con:
        save_dossier(con, d)
    return jsonify(d), 201


@app.get("/api/dossiers/<did>")
def dossier_get(did):
    denied = dossier_guard()
    if denied:
        return denied
    with _db_lock, db() as con:
        d = load_dossier(con, did)
    return jsonify(d) if d else error("Dossier introuvable.", 404)


@app.put("/api/dossiers/<did>")
def dossier_update(did):
    denied = dossier_guard()
    if denied:
        return denied
    data = request.get_json(silent=True) or {}
    with _db_lock, db() as con:
        d = load_dossier(con, did)
        if not d:
            return error("Dossier introuvable.", 404)
        for k, n in FIELDS.items():
            if k in data:
                d[k] = clip(data[k], n)
        if "pieces" in data:
            d["pieces"] = [str(x)[:40] for x in (data["pieces"] or [])][:40]
        clean_dossier(d)
        save_dossier(con, d)
    return jsonify(d)


@app.post("/api/dossiers/<did>/journal")
def dossier_journal_add(did):
    denied = dossier_guard()
    if denied:
        return denied
    data = request.get_json(silent=True) or {}
    texte = clip(data.get("texte"), 2000)
    if not texte:
        return error("Texte manquant.", 400)
    typ = data.get("type") if data.get("type") in ("rdv", "appel", "courrier", "demarche", "note") else "note"
    with _db_lock, db() as con:
        d = load_dossier(con, did)
        if not d:
            return error("Dossier introuvable.", 404)
        d["journal"].append({"id": uuid.uuid4().hex[:8], "type": typ, "date": clip(data.get("date"), 10),
                             "texte": texte, "auteur": clip(data.get("auteur"), 40)})
        save_dossier(con, d)
    return jsonify(d)


@app.delete("/api/dossiers/<did>/journal/<eid>")
def dossier_journal_delete(did, eid):
    denied = dossier_guard()
    if denied:
        return denied
    with _db_lock, db() as con:
        d = load_dossier(con, did)
        if not d:
            return error("Dossier introuvable.", 404)
        d["journal"] = [e for e in d["journal"] if e["id"] != eid]
        save_dossier(con, d)
    return jsonify(d)


@app.delete("/api/dossiers/<did>")
def dossier_delete(did):
    denied = dossier_guard()
    if denied:
        return denied
    with _db_lock, db() as con:
        con.execute("DELETE FROM dossiers WHERE id=?", (did,))
        con.commit()
        con.execute("VACUUM")  # efface réellement les données du fichier
    return jsonify({"ok": True})



# ---- export / import Excel
import csv
import datetime as _dt
import unicodedata

from openpyxl import Workbook, load_workbook

COLS = [("Prénom", "prenom"), ("Nom", "nom"), ("Langue", "langue_nom"), ("Téléphone", "tel"),
        ("Statut", "statut"), ("Échéance", "echeance"), ("Remarques", "notes"),
        ("Métier", "metier"), ("Niveau de français", "niveau_fr"), ("Disponibilité", "disponibilite"),
        ("Permis", "permis"), ("Entreprise", "entreprise"), ("Étape", "etape"), ("Type de titre", "titre_type"),
        ("Expiration du titre", "titre_expiration"), ("Consentement (date)", "consentement")]
PIECE_LABELS = {p["key"]: p["label"] for p in PARCOURS["pieces"]}
TYPE_LABEL = {"rdv": "Rendez-vous", "appel": "Appel", "demarche": "Démarche", "courrier": "Courrier", "note": "Note"}
LANG_KEYS = [  # (mot recherché, code navigateur, nom donné à Claude) — le plus précis d'abord
    ("algerien", "ar-DZ", "arabe algérien (darija)"), ("tunisien", "ar-TN", "arabe tunisien (darija)"),
    ("marocain", "ar-MA", "arabe marocain (darija)"), ("arabe", "ar-SA", "arabe standard"),
    ("anglais", "en-US", "anglais"), ("dari", "fa-AF", "dari (persan d'Afghanistan)"),
    ("pashto", "ps-AF", "pashto"), ("pachto", "ps-AF", "pashto"), ("turc", "tr-TR", "turc"),
    ("espagnol", "es-ES", "espagnol"), ("portugais", "pt-PT", "portugais"), ("russe", "ru-RU", "russe"),
    ("ukrainien", "uk-UA", "ukrainien"), ("bengali", "bn-BD", "bengali"), ("persan", "fa-IR", "persan"),
    ("farsi", "fa-IR", "persan"), ("ourdou", "ur-PK", "ourdou"), ("hindi", "hi-IN", "hindi"),
    ("chinois", "zh-CN", "chinois"), ("roumain", "ro-RO", "roumain"), ("italien", "it-IT", "italien"),
    ("allemand", "de-DE", "allemand"), ("polonais", "pl-PL", "polonais"), ("vietnamien", "vi-VN", "vietnamien"),
    ("albanais", "sq-AL", "albanais"), ("tamoul", "ta-IN", "tamoul"), ("somali", "so-SO", "somali"),
    ("amharique", "am-ET", "amharique"),
]
ALIASES = {
    "prenom": "prenom", "firstname": "prenom", "nom": "nom", "nomdefamille": "nom", "lastname": "nom",
    "langue": "langue", "language": "langue", "telephone": "tel", "tel": "tel", "phone": "tel", "portable": "tel",
    "statut": "statut", "status": "statut", "echeance": "echeance", "prochaineecheance": "echeance",
    "datelimite": "echeance", "remarques": "notes", "notes": "notes", "situation": "notes", "commentaire": "notes",
    "dossier": "dossier", "personne": "dossier", "date": "date", "type": "type", "note": "texte",
    "texte": "texte", "suivi": "texte", "auteur": "auteur", "benevole": "auteur",
    "metier": "metier", "niveaudefrancais": "niveau_fr", "niveaufrancais": "niveau_fr", "niveau": "niveau_fr",
    "disponibilite": "disponibilite", "permis": "permis", "entreprise": "entreprise", "employeur": "entreprise",
    "etape": "etape", "typedetitre": "titre_type", "titre": "titre_type", "expirationdutitre": "titre_expiration",
    "expirationtitre": "titre_expiration", "consentementdate": "consentement", "consentement": "consentement",
    "piecesrecues": "pieces",
}


def norm(v):
    t = unicodedata.normalize("NFD", str(v or "")).encode("ascii", "ignore").decode().lower()
    return "".join(c for c in t if c.isalnum())


def to_date(v):
    if isinstance(v, (_dt.datetime, _dt.date)):
        return v.strftime("%Y-%m-%d")
    t = str(v or "").strip()
    for f in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y", "%d/%m/%y"):
        try:
            return _dt.datetime.strptime(t[:10] if f == "%Y-%m-%d" else t, f).strftime("%Y-%m-%d")
        except ValueError:
            pass
    return ""


def lang_of(text):
    n = norm(text)
    for key, code, name in LANG_KEYS:
        if key in n:
            return code, name
    return "", str(text or "").strip()[:60]


def xlsx_response(wb, filename):
    buf = io.BytesIO()
    wb.save(buf)
    return Response(buf.getvalue(), mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="{filename}"', "Cache-Control": "no-store"})


def style_header(ws, widths):
    from openpyxl.styles import Font, PatternFill
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="0F5C6E")
    for i, w in enumerate(widths):
        ws.column_dimensions[chr(65 + i)].width = w
    ws.freeze_panes = "A2"


@app.get("/api/dossiers-export")
def dossiers_export():
    denied = dossier_guard()
    if denied:
        return denied
    with _db_lock, db() as con:
        ds = [d for d in (load_dossier(con, r[0]) for r in con.execute("SELECT id FROM dossiers")) if d]
    ds.sort(key=lambda d: (d.get("nom", "").lower(), d.get("prenom", "").lower()))
    wb = Workbook()
    ws = wb.active
    ws.title = "Dossiers"
    ws.append([h for h, _ in COLS] + ["Pièces reçues", "Pièces manquantes", "Créé le", "Dernière modification"])
    for d in ds:
        got = [k for k in PIECE_LABELS if k in (d.get("pieces") or [])]
        ws.append([d.get(k, "") for _, k in COLS] + [
            "; ".join(PIECE_LABELS[k] for k in got),
            "; ".join(v for k, v in PIECE_LABELS.items() if k not in got),
            _dt.datetime.fromtimestamp(d.get("created", 0)).strftime("%Y-%m-%d"),
            _dt.datetime.fromtimestamp(d.get("updated", 0)).strftime("%Y-%m-%d")])
    style_header(ws, [16, 18, 26, 16, 12, 12, 40, 22, 18, 22, 8, 24, 26, 22, 18, 18, 60, 60, 12, 20])
    sj = wb.create_sheet("Suivi")
    sj.append(["Dossier", "Date", "Type", "Note", "Auteur"])
    for d in ds:
        for e in d.get("journal", []):
            sj.append([f"{d.get('prenom', '')} {d.get('nom', '')}".strip(), e.get("date", ""),
                       TYPE_LABEL.get(e.get("type"), "Note"), e.get("texte", ""), e.get("auteur", "")])
    style_header(sj, [28, 12, 14, 70, 14])
    return xlsx_response(wb, f"dossiers-{_dt.date.today().isoformat()}.xlsx")


@app.get("/api/dossiers-modele")
def dossiers_modele():
    denied = dossier_guard()
    if denied:
        return denied
    wb = Workbook()
    ws = wb.active
    ws.title = "Dossiers"
    ws.append([h for h, _ in COLS])
    ex = {"prenom": "Amina", "nom": "Benali", "langue_nom": "arabe algérien", "tel": "06 00 00 00 00",
          "statut": "En cours", "echeance": "2026-12-01", "notes": "Exemple : à supprimer", "metier": "Aide-soignante",
          "niveau_fr": "Intermédiaire", "disponibilite": "Immédiate", "permis": "Non", "entreprise": "",
          "etape": PARCOURS["etapes"][1], "titre_type": "Récépissé", "titre_expiration": "2027-03-01",
          "consentement": "2026-10-01"}
    ws.append([ex.get(k, "") for _, k in COLS])
    style_header(ws, [16, 18, 26, 16, 12, 12, 40, 22, 18, 22, 8, 24, 26, 22, 18, 18])
    aide = wb.create_sheet("Aide")
    for line in ["Remplissez l'onglet « Dossiers » (une personne par ligne), puis importez le fichier dans Tarjam.",
                 "Colonnes : Prénom, Nom, Langue, Téléphone, Statut, Échéance, Remarques, Métier, Niveau de français (Débutant / Intermédiaire / Avancé), Disponibilité, Permis (Oui / Non), Entreprise, Étape, Type de titre, Expiration du titre, Consentement (date). Seul un prénom OU un nom est obligatoire.",
                 "Étape : " + ", ".join(PARCOURS["etapes"]) + ".",
                 "Pièces reçues (colonne facultative) : noms séparés par des points-virgules : " + " ; ".join(PIECE_LABELS.values()) + ".",
                 "Langue : arabe algérien, arabe tunisien, arabe marocain, arabe, anglais, dari, pashto, turc, espagnol, portugais, russe, ukrainien, bengali, persan, ourdou, hindi, chinois, roumain, italien, allemand, polonais, vietnamien, albanais, tamoul, somali, amharique.",
                 "Statut : En cours, En attente ou Clos (« En cours » par défaut).",
                 "Échéance : une date (2026-12-01 ou 01/12/2026).",
                 "Onglet facultatif « Suivi » : colonnes Dossier (Prénom Nom), Date, Type, Note, Auteur.",
                 "Les personnes déjà présentes (même prénom et même nom) sont ignorées."]:
        aide.append([line])
    aide.column_dimensions["A"].width = 130
    return xlsx_response(wb, "modele-dossiers.xlsx")


def read_rows(file):
    """Renvoie {nom_onglet: [dict colonne_normalisée -> valeur]} depuis un .xlsx ou un .csv."""
    name = (file.filename or "").lower()
    raw = file.read()
    if len(raw) > 5_000_000:
        raise ValueError("Fichier trop gros (5 Mo maximum).")
    sheets = {}
    if name.endswith(".csv"):
        text = raw.decode("utf-8-sig", errors="replace")
        dialect = csv.Sniffer().sniff(text[:2000], delimiters=";,\t") if text.strip() else csv.excel
        sheets["Dossiers"] = [list(r) for r in csv.reader(io.StringIO(text), dialect)]
    else:
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        for ws in wb.worksheets:
            sheets[ws.title] = [list(r) for r in ws.iter_rows(values_only=True)]
    out = {}
    for title, rows in sheets.items():
        rows = [r for r in rows if any(c not in (None, "") for c in r)][:2001]
        if not rows:
            continue
        keys = [ALIASES.get(norm(h), "") for h in rows[0]]
        out[title] = [{k: c for k, c in zip(keys, r) if k} for r in rows[1:]]
    return out


@app.post("/api/dossiers-import")
def dossiers_import():
    denied = dossier_guard()
    if denied:
        return denied
    f = request.files.get("file")
    if not f:
        return error("Aucun fichier reçu.", 400)
    try:
        sheets = read_rows(f)
    except Exception:
        return error("Fichier illisible : utilisez un .xlsx (Excel) ou .csv, de préférence le modèle.", 400)
    main = sheets.get("Dossiers") or next((v for k, v in sheets.items() if k != "Suivi"), [])
    suivi = sheets.get("Suivi", [])
    created, skipped, problems = 0, 0, []
    with _db_lock, db() as con:
        existing = {}
        for r in con.execute("SELECT id FROM dossiers"):
            d = load_dossier(con, r[0])
            if d:
                existing[norm(d.get("prenom")) + "|" + norm(d.get("nom"))] = d
        for i, row in enumerate(main, start=2):
            prenom, nom = clip(row.get("prenom"), 60), clip(row.get("nom"), 60)
            if not prenom and not nom:
                problems.append(f"Ligne {i} : ni prénom ni nom, ignorée.")
                continue
            key = norm(prenom) + "|" + norm(nom)
            if key in existing:
                skipped += 1
                continue
            code, lname = lang_of(row.get("langue"))
            statut = {"encours": "En cours", "enattente": "En attente", "clos": "Clos"}.get(norm(row.get("statut")), "En cours")
            if row.get("langue") and not code:
                problems.append(f"Ligne {i} : langue « {lname} » non reconnue (gardée telle quelle).")
            d = {k: clip(row.get(k), n) for k, n in FIELDS.items() if k in row}
            for k in FIELDS:
                d.setdefault(k, "")
            asked = {norm(x) for x in str(row.get("pieces") or "").replace(",", ";").split(";") if x.strip()}
            d.update({"prenom": prenom, "nom": nom, "langue_code": code, "langue_nom": lname, "statut": statut,
                      "pieces": [k for k, lab in PIECE_LABELS.items() if norm(lab) in asked or k in asked],
                      "created": time.time(), "journal": [], "id": uuid.uuid4().hex[:12]})
            if d["etape"]:
                d["etape"] = next((e for e in PARCOURS["etapes"] if norm(e) == norm(d["etape"])), d["etape"])
            clean_dossier(d)
            existing[key] = d
            d["_new"] = True
            created += 1
        for e in suivi:
            d = existing.get(norm(" ".join(str(e.get("dossier", "")).split()[:1])) + "|" + norm(" ".join(str(e.get("dossier", "")).split()[1:])))
            texte = clip(e.get("texte"), 2000)
            if not d or not d.get("_new") or not texte:
                continue
            t = norm(e.get("type"))
            typ = "rdv" if "rendez" in t else "appel" if "appel" in t else "demarche" if "demarche" in t else "courrier" if "courrier" in t else "note"
            d["journal"].append({"id": uuid.uuid4().hex[:8], "type": typ, "date": to_date(e.get("date")),
                                 "texte": texte, "auteur": clip(e.get("auteur"), 40)})
        for d in existing.values():
            if d.pop("_new", False):
                save_dossier(con, d)
    return jsonify({"crees": created, "ignores": skipped, "problemes": problems[:20]})

@app.get("/api/letters")
def letters():
    """Courriers standards (letters.json) + coordonnées de l'association."""
    denied = check_access()
    if denied:
        return denied
    with open(os.path.join(BASE_DIR, "letters.json"), encoding="utf-8") as f:
        templates = json.load(f)["templates"]
    return jsonify({
        "templates": templates,
        "asso": {
            "nom": os.environ.get("ASSO_NAME", "Nom de l'association"),
            "adresse": os.environ.get("ASSO_ADDRESS", "Adresse de l'association"),
            "contact": os.environ.get("ASSO_CONTACT", "Téléphone / e-mail"),
            "ville": os.environ.get("ASSO_CITY", ""),
        },
    })


@app.get("/dossiers.js")
def dossiers_js():
    return send_from_directory(STATIC_DIR, "dossiers.js", mimetype="text/javascript", max_age=0)


# ------------------------------------------------------------ assistant d'aide et notice
with open(os.path.join(BASE_DIR, "aide.md"), encoding="utf-8") as _f:
    HELP_KB = _f.read()


@app.post("/api/help")
def help_chat():
    """Assistant d'aide : répond aux questions sur l'utilisation de Tarjam (guide aide.md)."""
    denied = check_access()
    if denied:
        return denied
    if not os.environ.get("ANTHROPIC_API_KEY"):
        return error("L'assistant n'est pas disponible : clé API manquante sur le serveur.", 503)
    d = request.get_json(silent=True) or {}
    question = str(d.get("question", "")).strip()[:600]
    if not question:
        return error("Question vide.", 400)
    screen = str(d.get("screen", ""))[:30]
    msgs = []
    for m in (d.get("history") or [])[-6:]:
        if isinstance(m, dict) and m.get("role") in ("user", "assistant") and str(m.get("content", "")).strip():
            msgs.append({"role": m["role"], "content": str(m["content"])[:1200]})
    while msgs and msgs[0]["role"] != "user":
        msgs.pop(0)
    if msgs and msgs[-1]["role"] == "user":
        msgs.pop()
    msgs.append({"role": "user", "content": question})
    system = (
        "Tu es l'assistant d'aide intégré à l'application Tarjam. Tu aides des débutants (bénévoles "
        "d'une association) à utiliser l'application. Réponds en français simple et chaleureux (ou dans "
        "la langue de la question), en phrases courtes, avec des étapes numérotées quand il faut "
        "agir, en nommant les boutons exactement comme dans le guide. Appuie-toi UNIQUEMENT sur le guide "
        "ci-dessous : si l'information n'y est pas, dis-le et conseille de demander au responsable de "
        "l'association. Ne donne aucun conseil juridique, médical ou administratif et n'invente aucune "
        "fonction. Reste bref (8 lignes maximum sauf demande de détail). "
        f"L'utilisateur est actuellement sur l'écran : {screen or 'accueil'}.\n\n=== GUIDE ===\n{HELP_KB}"
    )
    try:
        msg = get_client().messages.create(model=MODEL, max_tokens=700, system=system, messages=msgs)
    except anthropic.APIError as e:
        return error(f"Erreur : {getattr(e, 'message', e)}", 502)
    return jsonify({"answer": "".join(b.text for b in msg.content if b.type == "text").strip()})


@app.get("/notice")
def notice():
    return send_from_directory(STATIC_DIR, "notice.html", mimetype="text/html", max_age=0)


@app.get("/aide.png")
def aide_png():
    return send_from_directory(STATIC_DIR, "aide.png", mimetype="image/png", max_age=86400)


@app.get("/aide-face.png")
def aide_face_png():
    return send_from_directory(STATIC_DIR, "aide-face.png", mimetype="image/png", max_age=86400)


@app.get("/aide.js")
def aide_js():
    return send_from_directory(STATIC_DIR, "aide.js", mimetype="text/javascript", max_age=0)

@app.get("/qr.svg")
def qr():
    text = request.args.get("text", "")[:300]
    if not text:
        return error("Texte manquant.", 400)
    buf = io.BytesIO()
    segno.make(text, error="m").save(buf, kind="svg", scale=6, border=2, dark="#0f5c6e")
    return Response(buf.getvalue(), mimetype="image/svg+xml", headers={"Cache-Control": "no-store"})


TASKS = {
    "traduire": (
        "Traduis fidèlement tout le texte lisible vers : {dst}, en gardant la structure (titres, listes, "
        "montants, dates). Mets entre [crochets] ce qui est illisible. "
        "Termine par une ligne « ➜ En bref : » suivie de 1 à 3 phrases très simples, en {dst}, qui disent "
        "de quoi il s'agit et ce que la personne doit faire (et avant quelle date, s'il y en a une)."
    ),
    "formulaire": (
        "C'est un formulaire à remplir. En {dst}, explique-le pour une personne qui ne lit pas bien le "
        "français : à quoi sert ce formulaire, puis chaque champ ou case dans l'ordre (nom du champ en "
        "français entre guillemets, ce qu'il faut y écrire en mots simples, un exemple). "
        "Indique les pièces à joindre, où signer et la date limite s'il y en a."
    ),
    "simplifier": (
        "Réécris ce document en langage très simple (niveau débutant, phrases courtes) en {dst}, avec "
        "ces rubriques : 1) De quoi s'agit-il ? 2) Ce que vous devez faire 3) Avant quelle date "
        "4) Ce qui se passe si vous ne faites rien. Pas de jargon : explique les mots administratifs."
    ),
}
DOC_MIMES = ("image/jpeg", "image/png", "image/webp", "image/gif")


@app.post("/api/document")
def document():
    denied = check_access()
    if denied:
        return denied
    if not os.environ.get("ANTHROPIC_API_KEY"):
        return error("Le serveur n'a pas encore de clé API Anthropic.", 503)
    f = request.files.get("image") or request.files.get("pdf")
    pasted = request.form.get("text", "").strip()[:20000]
    dst = request.form.get("to", "").strip()[:60]
    task = request.form.get("task", "traduire")
    if task not in TASKS or not dst or not (f or pasted):
        return error("Document, texte ou langue manquant.", 400)
    system = (
        "Tu es un traducteur dans une association qui aide des personnes étrangères. "
        "On te donne un document (courrier, formulaire, ordonnance, facture…). "
        + TASKS[task].format(dst=dst)
        + " Ajoute un pictogramme (emoji) au début des points importants pour faciliter la "
        "compréhension : 📅 date, 💶 argent, 📄 papier à fournir, 📍 lieu, 📞 téléphone, ⚠️ urgent, ✍️ signer. "
        "Si la langue demandée est une darija, écris en lettres arabes, dans le dialecte demandé, "
        "en gardant les termes administratifs français usuels (CAF, préfecture…). "
        "Ne rajoute aucun autre commentaire."
    )
    if pasted and not f:
        content = [{"type": "text", "text": "Voici le texte du document :\n\n" + pasted}]
    else:
        data = base64.standard_b64encode(f.read()).decode()
        if f.mimetype == "application/pdf" or (f.filename or "").lower().endswith(".pdf"):
            block = {"type": "document", "source": {"type": "base64", "media_type": "application/pdf", "data": data}}
        else:
            mime = f.mimetype if f.mimetype in DOC_MIMES else "image/jpeg"
            block = {"type": "image", "source": {"type": "base64", "media_type": mime, "data": data}}
        content = [block, {"type": "text", "text": "Voici le document."}]
    try:
        msg = get_client().messages.create(
            model=MODEL, max_tokens=6000, system=system,
            messages=[{"role": "user", "content": content}],
        )
    except anthropic.APIError as e:
        return error(f"Erreur de traduction : {getattr(e, 'message', e)}", 502)
    out = "".join(b.text for b in msg.content if b.type == "text").strip()
    return jsonify({"translation": out})


@app.post("/api/transcribe")
def transcribe():
    denied = check_access()
    if denied:
        return denied
    key = os.environ.get("OPENAI_API_KEY")
    if not key:
        return error("L'import de fichiers audio demande une clé OPENAI_API_KEY sur le serveur.", 503)
    f = request.files.get("audio")
    if not f:
        return error("Aucun fichier reçu.", 400)
    name = f.filename or "audio.ogg"
    if name.lower().endswith(".opus"):  # WhatsApp : même format qu'un .ogg
        name = name[:-5] + ".ogg"
    try:
        r = requests.post(
            "https://api.openai.com/v1/audio/transcriptions",
            headers={"Authorization": f"Bearer {key}"},
            data={"model": "whisper-1", "response_format": "verbose_json"},
            files={"file": (name, f.stream, f.mimetype or "application/octet-stream")},
            timeout=300,
        )
    except requests.RequestException as e:
        return error(f"Transcription impossible : {e}", 502)
    if r.status_code != 200:
        return error(f"Transcription refusée ({r.status_code}).", 502)
    j = r.json()
    return jsonify({"text": (j.get("text") or "").strip(), "language": j.get("language", "")})


if __name__ == "__main__":
    app.run(host="0.0.0.0", port=int(os.environ.get("PORT", "5000")))
